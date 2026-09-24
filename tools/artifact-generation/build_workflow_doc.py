from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REPO_ROOT = Path(__file__).resolve().parents[2]
OUTPUT = REPO_ROOT / "Deliverables" / "documents" / "Simple Clean End to End Workflow.docx"

NAVY = "1F4E78"
LIGHT_BLUE = "EAF2F8"
PALE_BLUE = "F5F9FC"
LIGHT_GRAY = "D9D9D9"
MID_GRAY = "5B6573"
BLACK = "000000"
WHITE = "FFFFFF"


def set_font(run, name="Aptos", size=None, bold=None, color=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=LIGHT_GRAY, size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_font(run, size=9, color=MID_GRAY)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.08
    p.add_run(text)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.08
    p.add_run(text)
    return p


def add_label_paragraph(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(label + " ")
    set_font(r, bold=True)
    p.add_run(text)
    return p


def add_stage(doc, number, title, purpose, records, handoff, bullets=None):
    heading = doc.add_heading(f"{number}  {title}", level=2)
    heading.paragraph_format.keep_with_next = True
    add_label_paragraph(doc, "Purpose", purpose)
    add_label_paragraph(doc, "What the system keeps", records)
    if bullets:
        for item in bullets:
            add_bullet(doc, item)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.keep_with_next = False
    r = p.add_run("Handoff to the next step  ")
    set_font(r, bold=True, color=NAVY)
    p.add_run(handoff)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, header in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        set_font(r, bold=True, color=WHITE, size=9.5)
        shade_cell(cell, NAVY)
        set_cell_margins(cell)
        set_cell_borders(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            cell.width = Inches(widths[i])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for i, value in enumerate(values):
            cell = row.cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            set_font(r, size=9.2)
            shade_cell(cell, PALE_BLUE if row_index % 2 else WHITE)
            set_cell_margins(cell)
            set_cell_borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cell.width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.68)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)
section.header_distance = Inches(0.3)
section.footer_distance = Inches(0.3)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.12

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
title_style.font.size = Pt(28)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor.from_string(BLACK)
title_style.paragraph_format.space_after = Pt(8)

for style_name, size, before, after in (("Heading 1", 17, 14, 7), ("Heading 2", 12.5, 10, 4), ("Heading 3", 11, 8, 3)):
    st = styles[style_name]
    st.font.name = "Aptos Display"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(BLACK)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for list_style in ("List Bullet", "List Bullet 2", "List Number"):
    styles[list_style].font.name = "Aptos"
    styles[list_style]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles[list_style]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles[list_style].font.size = Pt(10.2)

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = hp.add_run("Simple Clean End to End Workflow")
set_font(hr, size=8.5, color=MID_GRAY)
add_page_number(section.footer.paragraphs[0])

# Opening page
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(34)
p.paragraph_format.space_after = Pt(0)
r = p.add_run("SIMPLE CLEAN")
set_font(r, size=10, bold=True, color=NAVY)

doc.add_paragraph("End to End Workflow", style="Title")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(16)
r = p.add_run("How equipment, information, people, and customer communication move through one connected process")
set_font(r, size=13, color=MID_GRAY)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(14)
r = p.add_run("Prepared for William")
set_font(r, size=10.5, bold=True)
p.add_run("\nSeptember 2026")

doc.add_heading("The main idea", level=1)
p = doc.add_paragraph()
p.add_run("We will build one core tool with different screens for William, warehouse workers, technicians, sellers, buyers, and drivers. ").bold = True
p.add_run("Each team completes its part of the work, and the same machine record moves forward. Information should be entered once and reused for pricing, testing, listings, shipping, accounting, and warranty support.")

doc.add_heading("The flow in one view", level=1)
flow_rows = [
    ("1", "Acquire", "Seller inquiry, equipment photos, offer, pickup planning"),
    ("2", "Receive", "Unload, photograph, create machine record, attach QR, assign location"),
    ("3", "Prepare", "Inspect, test, repair, retest, clean, and approve for shipment"),
    ("4", "Sell", "List, answer inquiries, quote, collect payment, and reserve equipment"),
    ("5", "Deliver", "Pack, choose freight, dispatch, track, and confirm delivery"),
    ("6", "Support", "Handle warranty, exceptions, accounting, and management reporting"),
]
add_table(doc, ["Stage", "Area", "What happens"], flow_rows, widths=[0.55, 1.15, 5.2])

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(6)
r = p.add_run("Recommended build priority  ")
set_font(r, bold=True, color=NAVY)
p.add_run("Start with intake and machine identity, then test/repair/clean, then sales. Build sourcing automation after the warehouse process is reliable.")

doc.add_page_break()

# Connection map
doc.add_heading("How the components connect", level=1)
p = doc.add_paragraph("Every component produces a clear result that starts the next component. If something is missing or fails, the item goes to an exception queue instead of disappearing from the process.")

connection_rows = [
    ("1", "Seller inquiry", "Equipment information and photos", "Offer calculation"),
    ("2", "Offer and approval", "Accepted purchase terms", "Pickup planning"),
    ("3", "Pickup and inbound load", "Expected load and arrival", "Warehouse intake"),
    ("4", "Intake and QR", "One verified machine record", "Inspection and inventory"),
    ("5", "Inspection and production", "Passed, repair needed, or parts-only decision", "Listing or parts workflow"),
    ("6", "Parts", "Available/ordered part and recorded cost", "Repair and retest"),
    ("7", "Inventory and listings", "Available equipment published for sale", "Buyer inquiry"),
    ("8", "CRM and communication", "Qualified need and next action", "Quote or follow-up"),
    ("9", "Quote and payment", "Reserved equipment and sales order", "Production priority and packing"),
    ("10", "Packing and shipping", "Dispatched shipment with tracking", "Delivery"),
    ("11", "Delivery and warranty", "Proof of delivery or resolved claim", "Closed sale"),
    ("12", "Accounting and reporting", "Accurate cost, revenue, margin, and workload", "Management decisions"),
]
add_table(doc, ["#", "Component", "Result", "Connects to"], connection_rows, widths=[0.35, 1.65, 2.75, 2.25])

doc.add_heading("One machine record follows the entire journey", level=1)
add_numbered(doc, "The application creates an internal Machine ID as soon as the machine is received.")
add_numbered(doc, "A QR label is printed and attached to the machine.")
add_numbered(doc, "Workers scan the QR before testing, repairing, cleaning, moving, packing, or loading it.")
add_numbered(doc, "The nameplate photo, serial number, work history, parts, costs, sales order, shipment, and warranty all remain attached to that same record.")

doc.add_page_break()

# Phase 1
doc.add_heading("Phase 1  Acquire the equipment", level=1)

add_stage(
    doc, "1", "Seller or distributor inquiry",
    "Give a seller one clear way to describe the equipment they want Simple Clean to purchase.",
    "Seller contact, pickup address, access details, machine count, model and serial information, nameplate photos, condition photos, requested timing, and requested amount.",
    "The pricing component receives enough information to estimate resale value and inbound freight.",
    [
        "The Seller Portal is part of the same platform, but it has a simple external screen.",
        "Calls, texts, and emails can still be entered by staff when a seller does not use the portal.",
        "A nameplate photo helps identify model, serial, voltage, and phase. It does not prove condition or operation.",
    ],
)

add_stage(
    doc, "2", "Offer calculation and owner approval",
    "Create a consistent suggested offer while keeping William in control.",
    "Expected resale values, acquisition percentage, condition and configuration adjustments, estimated freight, packing/loading allowance, risk reserve, and each version of the offer.",
    "An accepted offer creates the expected inbound load and starts pickup preparation.",
    [
        "The starting acquisition range discussed was generally 25% to 30% of expected retail.",
        "Inbound freight lowers what Simple Clean can pay for the equipment.",
        "William reviews and approves the offer initially; the system should not send a binding offer automatically.",
    ],
)

add_stage(
    doc, "3", "Seller preparation and pickup",
    "Make sure the equipment is protected, loaded correctly, and expected by the warehouse.",
    "Packing instructions, preparation photos, pickup appointment, truck information, negotiated loading allowance, possible payment holdback, driver contact, and load status.",
    "The warehouse sees what is arriving, when it should arrive, and which machines are expected.",
    [
        "The seller or distributor normally packs and loads the equipment.",
        "Simple Clean may provide instructions, diagrams, or protective supplies.",
        "Packing/loading payments and any 10% to 15% damage holdback remain configurable terms.",
    ],
)

doc.add_page_break()

# Phase 2
doc.add_heading("Phase 2  Receive and prepare the equipment", level=1)

add_stage(
    doc, "4", "Warehouse intake and QR identity",
    "Create one trustworthy record for every physical machine without delaying the truck.",
    "Internal Machine ID, QR label, Acquisition Load, arrival photos, nameplate photo, confirmed make/model/serial, electrical information, location, and initial condition.",
    "The machine becomes visible in inventory and enters preliminary inspection.",
    [
        "Create the provisional Machine record first, then print its QR label.",
        "Scan the QR and attach the nameplate photo to that record.",
        "Computer vision suggests the plate information; a worker confirms or corrects it.",
        "If the plate is inaccessible, mark Nameplate Verification Required and continue receiving.",
    ],
)

add_stage(
    doc, "5", "Load costs and preliminary inspection",
    "Understand what each machine costs and quickly identify machines that should not consume repair time.",
    "Seller payment, inbound freight, packing/loading cost, allocation method, bearing/condition check, and repair-versus-parts decision.",
    "Repairable equipment enters production; uneconomical equipment enters the parts workflow.",
    [
        "Purchase cost can be allocated according to each machine's share of expected retail value.",
        "Freight can begin as a simple per-piece allocation and later use footprint or weight.",
        "A preliminary pass may allow an early listing, but the machine cannot ship before final QA.",
    ],
)

add_stage(
    doc, "6", "Test repair retest clean and QA",
    "Give workers a clear queue and create proof that every machine was prepared correctly.",
    "Assigned worker, checklist version, timestamps, test results, defects, repairs, parts, notes, required photos/videos, cleaning result, retest, and QA approval.",
    "A QA-released machine can be packed for an order; a blocked machine shows exactly what it needs.",
    [
        "Workers scan the QR before starting.",
        "Sold or reserved machines receive higher production priority.",
        "A failed test creates a defect. Repair and retest do not erase the original result.",
        "The washer and dryer paper checklists become controlled digital checklists.",
    ],
)

add_stage(
    doc, "7", "Parts requests and parts inventory",
    "Prevent machines from being forgotten when a required part is unavailable.",
    "Part needed, compatible machines, stock quantity, reservation, request status, preferred source, order batch, receipt, cost, and installation result.",
    "An available or received part returns the machine to repair and retest.",
    [
        "Technicians submit the need from the machine record.",
        "William receives one central weekly ordering queue instead of scattered notes.",
        "Receipts allocate the correct parts cost back to the machine.",
        "Parts recovered from scrapped machines enter inventory; surplus above a minimum level can be sold.",
    ],
)

doc.add_page_break()

# Phase 3
doc.add_heading("Phase 3  Market and sell the equipment", level=1)

add_stage(
    doc, "8", "Inventory listings and sales channels",
    "Create the listing once and reuse accurate information across sales channels.",
    "Price, available quantity, specifications, condition, photo set, description, approval, and publication status for Shopify, Facebook, and selected eBay listings.",
    "A buyer message or Shopify checkout enters the CRM and sales workflow.",
    [
        "Shopify becomes the digital showroom and checkout website.",
        "Facebook Marketplace remains an important lead source, but personal Marketplace posting stays supervised.",
        "A representative photo can be used only when it is clearly disclosed.",
        "Sold or reserved inventory must be removed or reduced across every channel.",
    ],
)

add_stage(
    doc, "9", "CRM communication and AI assistance",
    "Keep sales and service conversations together and make sure follow-ups are not missed.",
    "Customer or seller identity, calls, SMS, email, supported social messages, transcript/summary, related machine or opportunity, intent, promises, next action, and reply status.",
    "A qualified buyer need creates a quote; other conversations create the correct task, order update, shipment update, or claim.",
    [
        "Use a business email and local 602 business number so personal conversations are not collected.",
        "AI can classify, summarize, extract tasks, and draft SMS/email replies.",
        "William approves drafts initially. Calls are transcribed and summarized, not answered by a live AI voice agent.",
        "A daily dropped-ball list shows unanswered messages and overdue promises.",
    ],
)

add_stage(
    doc, "10", "Buyer wishlists and inventory alerts",
    "Remember what buyers need and contact them when matching equipment becomes available.",
    "Desired manufacturer/model/type/capacity, quantity, budget, location, timing, and separate email/SMS consent.",
    "A matching alert creates or reopens a sales opportunity when the buyer responds.",
    [
        "Customers can opt in through a Shopify popup, the portal, SMS confirmation, email, or staff-recorded verbal permission.",
        "Marketing permission is separate from normal order and shipment updates.",
        "Customers can opt out through STOP or an email unsubscribe link.",
    ],
)

add_stage(
    doc, "11", "Quote payment and reservation",
    "Turn interest into a clear order without selling the same equipment twice.",
    "Equipment, package discount, estimated freight, tax, delivery/access requirements, deposit, balance, expiration, accepted terms, reserved quantity, and exact machines when assigned.",
    "Payment creates or confirms the reservation, prioritizes required production, and creates the warehouse packing task.",
    [
        "The listed price does not replace a quote when freight, quantity, tax, or services vary.",
        "Inventory may be temporarily held with an expiration before payment.",
        "After the configured payment condition, the machine becomes Sold Awaiting Fulfillment.",
    ],
)

doc.add_page_break()

# Phase 4
doc.add_heading("Phase 4  Pack ship and deliver", level=1)

add_stage(
    doc, "12", "Warehouse packing and load planning",
    "Prepare the exact sold machines safely and record final shipment facts.",
    "Packing task, worker, scanned machines, QA verification, pallet assignment, protection checklist, final dimensions/weight, photos, and exceptions.",
    "A completed task marks the order Ready for Dispatch and provides accurate facts for booking.",
    [
        "The system creates a virtual pallet plan before the quote, then the warehouse confirms actual measurements after packing.",
        "Model rules determine units per pallet and protection needs.",
        "Warehouse workers own packing; technicians join when technical verification is needed.",
    ],
)

add_stage(
    doc, "13", "Delivery method carrier booking and tracking",
    "Choose the practical shipping method, book it, and keep the customer prepared.",
    "Parcel/LTL/dedicated/pickup option, carrier quote, accessorials, BOL, labels, pickup scan, security seal where used, tracking, arrival alerts, and proof of delivery.",
    "Proof of delivery starts the warranty period and closes fulfillment unless a delivery exception is opened.",
    [
        "Parts normally use parcel shipping. Small equipment orders usually use LTL.",
        "Medium orders may use a 26-foot dedicated truck with a liftgate.",
        "Large 53-foot loads normally require a dock or forklift.",
        "Arrival alerts help the customer prepare labor, access, dock, or forklift and reduce driver waiting time.",
    ],
)

doc.add_page_break()

# Phase 5
doc.add_heading("Phase 5  Support the customer and manage the business", level=1)

add_stage(
    doc, "14", "Warranty returns and exceptions",
    "Resolve problems with the full machine history available.",
    "Delivery date, 90-day warranty window, symptoms, evidence, test/repair/QA history, diagnosis, approved remedy, communication, parts, refund/replacement, and final cost.",
    "A resolved claim updates profitability and closes with a documented outcome.",
    [
        "The same exception process handles inbound damage, missing machines, shipment damage, returns, refunds, lost freight, or failed integrations.",
        "Critical issues alert William; routine issues enter the correct team's queue.",
    ],
)

add_stage(
    doc, "15", "Accounting dashboards and administration",
    "Give William accurate financial and operational visibility without turning the operational tool into accounting software.",
    "QuickBooks links, customer invoices/payments, vendor bills/expenses, refunds, approved cost entries, sync status, inventory aging, throughput, response time, margin, freight variance, and warranty cost.",
    "The dashboard helps William change priorities, prices, staffing, stock levels, and business rules.",
    [
        "The application keeps machine-level operating detail; QuickBooks remains the general ledger.",
        "Every dashboard number should open the machines, tasks, orders, or exceptions behind it.",
        "Admin settings control users, approvals, pricing modifiers, checklist versions, parts minimums, pallet rules, rates, warranty, notifications, and integrations.",
    ],
)

# People and screens
doc.add_heading("What each person sees", level=1)
role_rows = [
    ("William", "Approvals, sales pipeline, dropped-ball list, profitability, exceptions, settings"),
    ("Warehouse", "Arrivals, intake, locations, packing, loading, dispatch"),
    ("Technician", "Assigned machines, checklists, defects, repairs, parts, retests"),
    ("Cleaner", "Cleaning queue, checklist, photos, completion"),
    ("Seller", "Submission, offer, packing instructions, pickup status"),
    ("Buyer", "Shopify products, quote/order, payment, tracking, warranty request"),
    ("Driver", "Pickup/delivery details, location sharing, scans, proof"),
]
add_table(doc, ["Person", "Main view"], role_rows, widths=[1.3, 5.7])

doc.add_page_break()

# Build plan and discovery
doc.add_heading("Recommended implementation order", level=1)
build_rows = [
    ("1", "Foundation and existing inventory", "Machine record, separate statuses, spreadsheet cleanup/import, QR labels, users"),
    ("2", "Intake and production", "Tablet intake, location, inspection, testing, repair, parts, cleaning, QA"),
    ("3", "Sales control", "Listings, CRM, communications, quotes, reservations, Shopify, Facebook preparation"),
    ("4", "Fulfillment and support", "Packing, freight, tracking, delivery, warranty, returns, exceptions"),
    ("5", "Accounting and optimization", "QuickBooks, dashboards, wishlists, alerts, better pricing and load planning"),
    ("6", "Acquisition automation", "Seller Portal, automated draft offers, packing workflow, pickup tracking"),
]
add_table(doc, ["Order", "Delivery area", "What becomes usable"], build_rows, widths=[0.55, 2.0, 4.45])

doc.add_heading("The immediate next step", level=1)
p = doc.add_paragraph()
p.add_run("Visit the warehouse before finalizing the machine record or tablet workflow. ").bold = True
p.add_run("Observe real equipment moving through intake, testing, repair, waiting for parts, cleaning, QA, and packing. Interview workers separately and then review the combined workflow with William.")

doc.add_heading("What to confirm with William and the team", level=2)
for item in [
    "Who performs each step and who approves exceptions",
    "How workers prioritize and handle several machines at the same time",
    "Where machines wait at every stage and how locations are named",
    "What happens when a part is unavailable and how work resumes",
    "Exact pass/fail and required video rules for washer and dryer checklists",
    "Repair-versus-parts-only decisions",
    "QR label position, size, material, and printer location",
    "Exact pallet rules, truck limits, warranty terms, and approval limits",
    "Current business email, phone, Shopify, Facebook, eBay, freight, QuickBooks, and AI Surfer setup",
]:
    add_bullet(doc, item)

doc.add_heading("Definition of success", level=1)
p = doc.add_paragraph("The workflow is successful when William can open one system and answer four questions without chasing people or spreadsheets:")
for question in [
    "Where is every machine?",
    "What work has been completed and what is blocking it?",
    "What has been promised to a seller or buyer, and what needs attention next?",
    "What did each machine and load actually cost, sell for, and earn?",
]:
    add_bullet(doc, question)

# Document properties
doc.core_properties.title = "Simple Clean End to End Workflow"
doc.core_properties.subject = "Business workflow from equipment acquisition through delivery and warranty"
doc.core_properties.author = "Simple Clean"
doc.core_properties.keywords = "workflow, inventory, warehouse, sales, shipping, warranty"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
